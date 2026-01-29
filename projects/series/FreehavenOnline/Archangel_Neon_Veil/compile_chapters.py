import os
from docx import Document
from docx.shared import Pt

def compile_chapters():
    base_path = "projects/series/FreehavenOnline/Archangel_Neon_Veil"
    manuscript_path = os.path.join(base_path, "manuscript")
    output_path = os.path.join(base_path, "output", "doc", "Archangel_Neon_Veil_Chapters_1-4.docx")

    doc = Document()

    chapters = ["Chapter_1.md", "Chapter_2.md", "Chapter_3.md", "Chapter_4.md"]

    for chapter_file in chapters:
        file_path = os.path.join(manuscript_path, chapter_file)
        if os.path.exists(file_path):
            with open(file_path, "r", encoding="utf-8") as f:
                lines = f.readlines()

            for line in lines:
                line = line.strip()
                if not line:
                    continue

                if line.startswith("# "):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith("### "):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith("***"):
                    doc.add_paragraph("***", style='Normal').alignment = 1 # Center
                else:
                    doc.add_paragraph(line)

            doc.add_page_break()
        else:
            print(f"Warning: {chapter_file} not found.")

    doc.save(output_path)
    print(f"Saved compiled manuscript to {output_path}")

if __name__ == "__main__":
    compile_chapters()
