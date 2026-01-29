import os
import shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
import markdown2
import re
import argparse

# Configuration
AUTHOR = "Jun Prince"
PROJECT_NAME = "Neon_Veil"
OUTPUT_BASE = "output/Neon_Veil"
TITLE = "NEON VEIL"

def ensure_dirs():
    for fmt in ['doc', 'docx', 'pdf', 'txt']:
        os.makedirs(os.path.join(OUTPUT_BASE, fmt), exist_ok=True)

def read_markdown(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        return f.read()

def create_txt(content, output_path):
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(content)
    print(f"Created TXT: {output_path}")

def create_docx(content, output_path, legacy_doc=False):
    doc = Document()

    # Title Page
    doc.add_heading(TITLE, 0)
    doc.add_paragraph(f"By {AUTHOR}")
    doc.add_paragraph(f"Word Count: {len(content.split())}")
    doc.add_page_break()

    lines = content.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue

        if line.startswith('#'):
            # Heading
            level = min(line.count('#'), 3)
            text = line.lstrip('#').strip()
            if "CHAPTER" in text or "SCENE" in text:
                doc.add_heading(text, level=1)
            else:
                doc.add_heading(text, level=level)
        elif line.startswith('**') and line.endswith('**') and len(line) < 100:
             # Bold line / Subheader / HUD element
            p = doc.add_paragraph()
            run = p.add_run(line.replace('**', ''))
            run.bold = True
        else:
            # Standard paragraph
            p = doc.add_paragraph(line)
            p_fmt = p.paragraph_format
            p_fmt.first_line_indent = Inches(0.25)

    doc.save(output_path)
    print(f"Created DOCX: {output_path}")

    if legacy_doc:
        doc_path = output_path.replace('.docx', '.doc').replace('/docx/', '/doc/')
        shutil.copy(output_path, doc_path)
        print(f"Created DOC: {doc_path}")

def create_pdf(content, output_path):
    class PDF(FPDF):
        def header(self):
            self.set_font('Times', 'I', 8)
            self.cell(0, 10, f'{AUTHOR} / {TITLE}', new_x="LMARGIN", new_y="NEXT", align='R')

        def footer(self):
            self.set_y(-15)
            self.set_font('Times', 'I', 8)
            self.cell(0, 10, f'Page {self.page_no()}', align='C')

    pdf = PDF()
    pdf.add_page()

    # Title Page
    pdf.set_font('Times', 'B', 24)
    pdf.cell(0, 60, TITLE, new_x="LMARGIN", new_y="NEXT", align='C')
    pdf.set_font('Times', '', 14)
    pdf.cell(0, 10, f"By {AUTHOR}", new_x="LMARGIN", new_y="NEXT", align='C')
    pdf.add_page()

    # Body
    pdf.set_font('Times', '', 12)

    # Replace any potential non-latin characters that might crash the standard font
    sanitized_content = content.replace('—', '-').replace('“', '"').replace('”', '"').replace('’', "'").replace('…', '...')
    sanitized_content = sanitized_content.encode('latin-1', 'replace').decode('latin-1')

    pdf.set_left_margin(15)
    pdf.set_right_margin(15)

    lines = sanitized_content.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            pdf.ln(5)
            continue

        if "CHAPTER" in line or "SCENE" in line:
             pdf.set_font('Times', 'B', 14)
             pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT", align='L')
             pdf.set_font('Times', '', 12)
        else:
             # Split very long words that might break FPDF
             words = line.split()
             safe_words = []
             for w in words:
                 if len(w) > 50: # Arbitrary limit for a "word"
                     safe_words.append(w[:50] + "-")
                     safe_words.append(w[50:])
                 else:
                     safe_words.append(w)
             safe_line = " ".join(safe_words)

             try:
                # 0 width = use remaining space
                pdf.multi_cell(0, 6, safe_line, new_x="LMARGIN", new_y="NEXT", align='L')
             except Exception as e:
                print(f"Warning: Skipping problematic line in PDF: {safe_line[:20]}... ({e})")

    pdf.output(output_path)
    print(f"Created PDF: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Convert Markdown manuscript to DOCX, PDF, and TXT.")
    parser.add_argument("--file", help="Path to the source markdown file", required=True)
    args = parser.parse_args()

    ensure_dirs()
    content = read_markdown(args.file)

    base_name = os.path.splitext(os.path.basename(args.file))[0]
    # If the file is like Chapter_1.md, we want output like Neon_Veil_Chapter_1.docx
    # Or just use the base name if it's unique enough. Let's use Project_Chapter format.
    out_name = f"{PROJECT_NAME}_{base_name}"

    # TXT
    create_txt(content, os.path.join(OUTPUT_BASE, 'txt', f'{out_name}.txt'))

    # DOCX and DOC
    create_docx(content, os.path.join(OUTPUT_BASE, 'docx', f'{out_name}.docx'), legacy_doc=True)

    # PDF
    create_pdf(content, os.path.join(OUTPUT_BASE, 'pdf', f'{out_name}.pdf'))

if __name__ == "__main__":
    main()
