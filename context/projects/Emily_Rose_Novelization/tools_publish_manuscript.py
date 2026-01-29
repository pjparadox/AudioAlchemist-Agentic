from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
import os

def create_manuscript():
    # Merge files
    files = [
        "context/projects/Emily_Rose_Novelization/output_quill/Chapter_1_3.txt",
        "context/projects/Emily_Rose_Novelization/output_quill/Chapter_4_6.txt",
        "context/projects/Emily_Rose_Novelization/output_quill/Chapter_7_8.txt",
        "context/projects/Emily_Rose_Novelization/output_quill/Chapter_9_11.txt"
    ]

    full_text = ""
    for f in files:
        if os.path.exists(f):
            with open(f, 'r', encoding='utf-8') as infile:
                full_text += infile.read() + "\n\n"
        else:
            print(f"Warning: {f} not found.")

    output_txt = "context/projects/Emily_Rose_Novelization/output_quill/The_Exorcism_of_Emily_Rose_Novelization_Complete.txt"
    with open(output_txt, 'w', encoding='utf-8') as out:
        out.write(full_text)

    # Create DOCX
    doc = Document()
    doc.add_heading('The Exorcism of Emily Rose', 0)
    doc.add_paragraph('A Novelization')
    doc.add_paragraph('By Jules (Agent Suite)')
    doc.add_page_break()

    # Add text to docx preserving basic formatting
    for line in full_text.split('\n'):
        if line.startswith('Chapter') or line.startswith('***'):
            doc.add_heading(line, level=1)
        else:
            doc.add_paragraph(line)

    docx_path = "context/projects/Emily_Rose_Novelization/output_quill/The_Exorcism_of_Emily_Rose_Complete.docx"
    doc.save(docx_path)
    print(f"Created {docx_path}")

    # Create PDF
    pdf_path = "context/projects/Emily_Rose_Novelization/output_quill/The_Exorcism_of_Emily_Rose_Complete.pdf"
    c = canvas.Canvas(pdf_path, pagesize=letter)
    width, height = letter

    c.setFont("Helvetica-Bold", 24)
    c.drawCentredString(width/2, height-100, "The Exorcism of Emily Rose")
    c.setFont("Helvetica", 12)
    c.drawCentredString(width/2, height-130, "A Novelization")
    c.showPage()

    text_object = c.beginText(40, height - 40)
    text_object.setFont("Helvetica", 10)

    for line in full_text.split('\n'):
        if text_object.getY() < 40:
            c.drawText(text_object)
            c.showPage()
            text_object = c.beginText(40, height - 40)
            text_object.setFont("Helvetica", 10)

        if line.startswith('Chapter'):
            text_object.setFont("Helvetica-Bold", 12)
            text_object.textLine(line)
            text_object.setFont("Helvetica", 10)
        else:
            # Simple word wrap logic for PDF
            words = line.split()
            line_buffer = ""
            for word in words:
                if c.stringWidth(line_buffer + " " + word, "Helvetica", 10) < (width - 80):
                    line_buffer += " " + word
                else:
                    text_object.textLine(line_buffer)
                    line_buffer = word
            text_object.textLine(line_buffer)

    c.drawText(text_object)
    c.save()
    print(f"Created {pdf_path}")

if __name__ == "__main__":
    create_manuscript()
