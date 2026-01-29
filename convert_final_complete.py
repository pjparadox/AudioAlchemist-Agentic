import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit

INPUT_FILE = "output_quill/Deadpool_Wolverine_Novelization_Final_Complete.txt"
OUTPUT_DOCX = "output_quill/Deadpool_Wolverine_Novelization_Final_Complete.docx"
OUTPUT_DOC = "output_quill/Deadpool_Wolverine_Novelization_Final_Complete.doc"
OUTPUT_PDF = "output_quill/Deadpool_Wolverine_Novelization_Final_Complete.pdf"

def create_docx(text, filepath):
    doc = Document()
    doc.add_heading('DEADPOOL & WOLVERINE: THE COMPLETE NOVELIZATION', 0)
    for line in text.split('\n'):
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('**') and line.endswith('**'):
             doc.add_heading(line.strip('*'), level=2)
        else:
            doc.add_paragraph(line)
    doc.save(filepath)
    print(f"Saved {filepath}")

def create_pdf(text, filepath):
    c = canvas.Canvas(filepath, pagesize=letter)
    width, height = letter
    y = height - 40
    c.setFont("Helvetica", 12)

    for line in text.split('\n'):
        if line.startswith('# '):
             c.setFont("Helvetica-Bold", 14)
             text_to_draw = line[2:]
        else:
             c.setFont("Helvetica", 12)
             text_to_draw = line

        # Simple wrapping
        lines = simpleSplit(text_to_draw, c._fontname, c._fontsize, width - 80)
        for subline in lines:
            if y < 40:
                c.showPage()
                y = height - 40
                if line.startswith('# '): c.setFont("Helvetica-Bold", 14)
                else: c.setFont("Helvetica", 12)

            c.drawString(40, y, subline)
            y -= 14
        y -= 6

    c.save()
    print(f"Saved {filepath}")

def main():
    if not os.path.exists(INPUT_FILE):
        print("Input file not found.")
        return

    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        content = f.read()

    # DOCX
    create_docx(content, OUTPUT_DOCX)

    # DOC (Renamed DOCX)
    import shutil
    shutil.copy(OUTPUT_DOCX, OUTPUT_DOC)
    print(f"Saved {OUTPUT_DOC}")

    # PDF
    create_pdf(content, OUTPUT_PDF)

if __name__ == "__main__":
    main()
