import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit

INPUT_FILES = [
    ("output_quill/Chapter_17_UltraDense.md", "DEADPOOL & WOLVERINE: CHAPTER 17 (ULTRA-DENSITY)"),
    ("output_quill/Chapter_18_UltraDense.md", "DEADPOOL & WOLVERINE: CHAPTER 18 (ULTRA-DENSITY)")
]

def create_docx(text, title, filepath):
    doc = Document()
    doc.add_heading(title, 0)
    for line in text.split('\n'):
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('**') and line.endswith('**'):
             doc.add_heading(line.strip('*'), level=2)
        else:
            doc.add_paragraph(line)
    doc.save(filepath)
    print(f"Saved {filepath}")

def create_pdf(text, filepath):
    c = canvas.Canvas(filepath, pagesize=letter)
    width, height = letter
    y = height - 40
    c.setFont("Helvetica", 12)

    for line in text.split('\n'):
        if line.startswith('# '):
             c.setFont("Helvetica-Bold", 14)
             text_to_draw = line[2:]
        else:
             c.setFont("Helvetica", 12)
             text_to_draw = line

        lines = simpleSplit(text_to_draw, c._fontname, c._fontsize, width - 80)
        for subline in lines:
            if y < 40:
                c.showPage()
                y = height - 40
                if line.startswith('# '): c.setFont("Helvetica-Bold", 14)
                else: c.setFont("Helvetica", 12)

            c.drawString(40, y, subline)
            y -= 14
        y -= 6

    c.save()
    print(f"Saved {filepath}")

def main():
    for input_path, title in INPUT_FILES:
        if not os.path.exists(input_path):
            print(f"Input file not found: {input_path}")
            continue

        with open(input_path, "r", encoding="utf-8") as f:
            content = f.read()

        base_name = os.path.splitext(os.path.basename(input_path))[0]
        output_docx = f"output_quill/Deadpool_Wolverine_{base_name}.docx"
        output_doc = f"output_quill/Deadpool_Wolverine_{base_name}.doc"
        output_pdf = f"output_quill/Deadpool_Wolverine_{base_name}.pdf"

        # DOCX
        create_docx(content, title, output_docx)

        # DOC (Renamed DOCX)
        import shutil
        shutil.copy(output_docx, output_doc)
        print(f"Saved {output_doc}")

        # PDF
        create_pdf(content, output_pdf)

if __name__ == "__main__":
    main()
