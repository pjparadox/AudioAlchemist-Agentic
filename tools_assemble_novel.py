import os
import shutil
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

OUTPUT_DIR = "output_quill"
PARTS = [
    "Prologue_The_Barn.txt",
    "Extreme_Chapter_1.txt",
    "Extreme_Chapter_2.txt",
    "Extreme_Chapter_3.txt",
    "Extreme_Chapter_4.txt",
    "Extreme_Chapter_5.txt",
    "Extreme_Chapter_6.txt",
    "Extreme_Chapter_7.txt",
    "Extreme_Chapter_8.txt",
    "Extreme_Chapter_9.txt",
    "Extreme_Chapter_10.txt"
]
FINAL_NAME = "The_Exorcism_of_Emily_Rose_Part_I_Complete_Ultra_Density"

def read_file(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()

def create_docx(text, filename):
    doc = Document()
    doc.add_heading("The Exorcism of Emily Rose", 0)
    doc.add_heading("Part I: The Body / The Case Opens (Ultra-Density)", 1)

    for line in text.split("\n"):
        if line.startswith("#"):
            level = line.count("#")
            clean_line = line.replace("#", "").strip()
            if level > 9: level = 9
            doc.add_heading(clean_line, level=level)
        elif line.strip() == "":
            pass
        else:
            doc.add_paragraph(line)
    doc.save(filename)
    print(f"Saved DOCX: {filename}")

def create_pdf(text, filename):
    doc = SimpleDocTemplate(filename, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("The Exorcism of Emily Rose", styles['Title']))
    story.append(Paragraph("Part I: The Body / The Case Opens (Ultra-Density)", styles['Heading1']))
    story.append(Spacer(1, 24))

    for line in text.split("\n"):
        if line.startswith("#"):
            level = line.count("#")
            clean_line = line.replace("#", "").strip()
            style = styles['Heading1'] if level == 1 else styles['Heading2']
            story.append(Paragraph(clean_line, style))
        elif line.strip():
            story.append(Paragraph(line, styles['Normal']))
        story.append(Spacer(1, 12))

    doc.build(story)
    print(f"Saved PDF: {filename}")

def main():
    print("Assembling Part I (Ultra-Density)...")
    full_text = ""
    for part in PARTS:
        path = os.path.join(OUTPUT_DIR, part)
        if os.path.exists(path):
            full_text += read_file(path) + "\n\n"
        else:
            print(f"Missing: {path}")

    # Save TXT
    txt_path = os.path.join(OUTPUT_DIR, FINAL_NAME + ".txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(full_text)
    print(f"Saved TXT: {txt_path}")

    # Save DOCX
    docx_path = os.path.join(OUTPUT_DIR, FINAL_NAME + ".docx")
    create_docx(full_text, docx_path)

    # Save DOC (Renamed DOCX)
    doc_path = os.path.join(OUTPUT_DIR, FINAL_NAME + ".doc")
    shutil.copy(docx_path, doc_path)
    print(f"Saved DOC: {doc_path}")

    # Save PDF
    pdf_path = os.path.join(OUTPUT_DIR, FINAL_NAME + ".pdf")
    create_pdf(full_text, pdf_path)

    word_count = len(full_text.split())
    print(f"Total Word Count: {word_count}")

if __name__ == "__main__":
    main()
